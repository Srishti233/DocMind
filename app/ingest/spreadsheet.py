"""CSV/XLSX ingestion into SQLite tables (rows are NEVER embedded).

A schema-description chunk is produced per table so the router/retriever can
discover spreadsheets; aggregation questions are answered with text-to-SQL.
"""
from __future__ import annotations

import csv
import json
import re
import sqlite3
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Iterator

from app import config
from app.db import connect
from app.ingest.parsers import ParseError

BATCH = 1000

_META = """
CREATE TABLE IF NOT EXISTS sheet_tables (
    table_name TEXT PRIMARY KEY,
    doc_id TEXT NOT NULL,
    workspace TEXT NOT NULL,
    filename TEXT NOT NULL,
    sheet TEXT,
    columns TEXT NOT NULL,
    row_count INTEGER NOT NULL,
    sample TEXT
);
"""


@dataclass
class TableInfo:
    """Metadata about one loaded table."""

    table_name: str
    doc_id: str
    workspace: str
    filename: str
    sheet: str | None
    columns: list[tuple[str, str]]  # (name, SQLite type)
    row_count: int
    sample: list[list[Any]] = field(default_factory=list)


def sanitize_ident(name: str, fallback: str = "col") -> str:
    """Lowercase snake_case identifier safe for SQL."""
    s = re.sub(r"[^0-9a-zA-Z]+", "_", str(name).strip()).strip("_").lower()
    if not s:
        s = fallback
    if s[0].isdigit():
        s = f"{fallback}_{s}"
    return s[:60]


def _unique_columns(header: Iterable[Any]) -> list[str]:
    seen: dict[str, int] = {}
    out: list[str] = []
    for i, h in enumerate(header):
        base = sanitize_ident(h if h is not None else "", f"col{i + 1}")
        n = seen.get(base, 0)
        seen[base] = n + 1
        out.append(base if n == 0 else f"{base}_{n + 1}")
    return out


def _coerce(v: Any) -> Any:
    """Turn CSV strings into int/float where possible; empty -> None."""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if s == "":
            return None
        try:
            return int(s)
        except ValueError:
            pass
        try:
            return float(s.replace(",", "")) if re.fullmatch(r"-?[\d,]*\.?\d+(?:[eE][-+]?\d+)?", s) else s
        except ValueError:
            return s
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return v


def _infer_types(rows: list[list[Any]], ncols: int) -> list[str]:
    types = []
    for c in range(ncols):
        vals = [r[c] for r in rows if c < len(r) and r[c] is not None]
        if vals and all(isinstance(v, int) and not isinstance(v, bool) for v in vals):
            types.append("INTEGER")
        elif vals and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals):
            types.append("REAL")
        else:
            types.append("TEXT")
    return types


def init_sheets_db() -> None:
    """Create the table-registry in the sheets database."""
    with connect(config.settings.sheets_db) as con:
        con.executescript(_META)


def _unique_table_name(con: sqlite3.Connection, base: str) -> str:
    taken = {r[0] for r in con.execute("SELECT table_name FROM sheet_tables")}
    taken |= {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    name, i = base, 2
    while name in taken:
        name, i = f"{base}_{i}", i + 1
    return name


def _load(con: sqlite3.Connection, doc: dict[str, Any], base: str, sheet: str | None,
          header: list[Any], rows: Iterator[list[Any]]) -> TableInfo | None:
    """Create a table from a header and a streaming row iterator."""
    cols = _unique_columns(header)
    first: list[list[Any]] = []
    for r in rows:
        first.append([_coerce(v) for v in r])
        if len(first) >= BATCH:
            break
    if not first:
        return None
    types = _infer_types(first, len(cols))
    table = _unique_table_name(con, base)
    con.execute(f'CREATE TABLE "{table}" (' + ", ".join(f'"{c}" {t}' for c, t in zip(cols, types)) + ")")
    ins = f'INSERT INTO "{table}" VALUES (' + ",".join("?" * len(cols)) + ")"

    def norm(r: list[Any]) -> list[Any]:
        r = list(r[:len(cols)]) + [None] * (len(cols) - len(r))
        return r

    total = 0
    batch = [norm(r) for r in first]
    while batch:
        con.executemany(ins, batch)
        total += len(batch)
        batch = []
        for r in rows:
            batch.append(norm([_coerce(v) for v in r]))
            if len(batch) >= BATCH:
                break
    info = TableInfo(table, doc["doc_id"], doc["workspace"], doc["filename"], sheet,
                     list(zip(cols, types)), total, [norm(r) for r in first[:3]])
    con.execute("INSERT INTO sheet_tables VALUES (?,?,?,?,?,?,?,?)",
                (table, info.doc_id, info.workspace, info.filename, sheet,
                 json.dumps(info.columns), total, json.dumps(info.sample, default=str)))
    return info


def _csv_rows(path: Path) -> tuple[list[str], Iterator[list[str]]]:
    f = open(path, "r", encoding="utf-8-sig", errors="replace", newline="")
    sample = f.read(4096)
    f.seek(0)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(f, dialect)
    try:
        header = next(reader)
    except StopIteration:
        f.close()
        raise ParseError("The CSV file is empty.") from None

    def gen() -> Iterator[list[str]]:
        try:
            for row in reader:
                if any(c.strip() for c in row):
                    yield row
        finally:
            f.close()

    return header, gen()


def load_file(doc: dict[str, Any]) -> list[TableInfo]:
    """Load a CSV/XLSX registry document into SQLite. Returns the created tables."""
    init_sheets_db()
    path = Path(doc["path"])
    ext = path.suffix.lower()
    stem = sanitize_ident(Path(doc["filename"]).stem, "sheet")
    infos: list[TableInfo] = []
    con = sqlite3.connect(str(config.settings.sheets_db), timeout=30)
    con.isolation_level = None  # manual transaction so CREATE TABLE is rolled back on failure
    con.execute("BEGIN")
    try:
        if ext == ".csv":
            header, rows = _csv_rows(path)
            info = _load(con, doc, stem, None, header, rows)
            if info:
                infos.append(info)
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            try:
                wb = load_workbook(str(path), read_only=True, data_only=True)
            except Exception as e:  # noqa: BLE001
                raise ParseError(f"Could not read the XLSX: {e}") from e
            try:
                multi = len(wb.sheetnames) > 1
                for ws in wb.worksheets:
                    it = ws.iter_rows(values_only=True)
                    header: list[Any] | None = None
                    for row in it:
                        if any(c is not None and str(c).strip() for c in row):
                            header = list(row)
                            break
                    if header is None:
                        continue
                    rows = (list(r) for r in it if any(c is not None and str(c).strip() for c in r))
                    base = f"{stem}_{sanitize_ident(ws.title, 'sheet')}" if multi else stem
                    info = _load(con, doc, base, ws.title, header, rows)
                    if info:
                        infos.append(info)
            finally:
                wb.close()
        else:
            raise ParseError(f"Not a spreadsheet: {ext}")
        if not infos:
            con.rollback()
            raise ParseError("The spreadsheet contains no data rows.")
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()
    return infos


def schema_text(info: TableInfo) -> str:
    """Natural-language schema description that is embedded for discovery."""
    cols = ", ".join(f"{n} ({t})" for n, t in info.columns)
    sheet = f", sheet '{info.sheet}'" if info.sheet else ""
    sample = "; ".join(", ".join(f"{c}={v}" for (c, _), v in zip(info.columns, row))
                       for row in info.sample[:2])
    return (f"Spreadsheet table `{info.table_name}` from file {info.filename}{sheet} with "
            f"{info.row_count} rows. Columns: {cols}. Example rows: {sample}.")


def drop_tables(doc_id: str) -> None:
    """Drop every table created for a document."""
    path = config.settings.sheets_db
    if not path.exists():
        return
    init_sheets_db()
    con = sqlite3.connect(str(path), timeout=30)
    try:
        for (t,) in con.execute("SELECT table_name FROM sheet_tables WHERE doc_id=?", (doc_id,)).fetchall():
            con.execute(f'DROP TABLE IF EXISTS "{t}"')
        con.execute("DELETE FROM sheet_tables WHERE doc_id=?", (doc_id,))
        con.commit()
    finally:
        con.close()


def list_tables(workspace: str, filename: str | None = None,
                doc_ids: list[str] | None = None) -> list[TableInfo]:
    """Tables available in a workspace (optionally limited to a file / docs)."""
    path = config.settings.sheets_db
    if not path.exists():
        return []
    init_sheets_db()
    q, args = "SELECT * FROM sheet_tables WHERE workspace=?", [workspace]
    if filename:
        q += " AND filename=?"
        args.append(filename)
    if doc_ids:
        q += f" AND doc_id IN ({','.join('?' * len(doc_ids))})"
        args.extend(doc_ids)
    with connect(path) as con:
        rows = con.execute(q + " ORDER BY table_name", args).fetchall()
    return [TableInfo(r["table_name"], r["doc_id"], r["workspace"], r["filename"], r["sheet"],
                      [tuple(c) for c in json.loads(r["columns"])], r["row_count"],
                      json.loads(r["sample"] or "[]")) for r in rows]
